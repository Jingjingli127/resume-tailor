import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / ".agents/skills/resume-tailor/scripts/retrieve-project-evidence.py"
SPEC = importlib.util.spec_from_file_location("retrieve_project_evidence", SCRIPT)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class RequirementParsingTests(unittest.TestCase):
    def test_structured_jd_ignores_title_and_weights_categories(self):
        blocks = [
            ("Example Company - Senior Data Scientist", False),
            ("Responsibilities", True),
            ("Build forecasting pipelines for product decisions.", False),
            ("Minimum Qualifications", True),
            ("Strong Python and SQL experience.", False),
            ("Preferred Qualifications", True),
            ("Advertising experience is a plus.", False),
        ]

        requirements = MODULE.requirement_records_from_blocks(blocks)

        self.assertEqual([item["requirement_id"] for item in requirements], ["R1", "R2", "R3"])
        self.assertEqual(
            [item["category"] for item in requirements],
            ["responsibility", "minimum_qualification", "preferred_qualification"],
        )
        self.assertGreater(requirements[0]["weight"], requirements[2]["weight"])
        self.assertNotIn("Example Company", " ".join(item["text"] for item in requirements))

    def test_unstructured_query_remains_supported(self):
        requirements = MODULE.requirements_from_text("Python SQL forecasting")
        self.assertEqual(len(requirements), 1)
        self.assertEqual(requirements[0]["category"], "other")

    def test_position_summary_is_role_context(self):
        requirements = MODULE.requirements_from_text(
            "Position Summary\nSupport operations through advanced analytics and scalable reporting.\n"
            "Job Responsibilities\nBuild automated dashboards."
        )

        self.assertEqual(
            [item["category"] for item in requirements],
            ["role_context", "responsibility"],
        )

    def test_short_required_and_preferred_headings_are_weighted(self):
        requirements = MODULE.requirements_from_text(
            "Workstreams & Tasks\nBuild reproducible data pipelines.\n"
            "Required\nUse Python and SQL.\n"
            "Preferred\nExperience with cloud platforms."
        )

        self.assertEqual(
            [item["category"] for item in requirements],
            ["responsibility", "minimum_qualification", "preferred_qualification"],
        )


class CandidateStrengthTests(unittest.TestCase):
    def classify(self, requirement, evidence):
        return MODULE.candidate_strength(requirement, evidence, {}, 10)

    def test_direct_evidence_requires_substantial_explicit_coverage(self):
        result = self.classify(
            "Build Python SQL forecasting pipelines",
            "Built Python SQL forecasting pipelines for product planning",
        )
        self.assertEqual(result["classification"], "direct")

    def test_adjacent_competency_is_transferable(self):
        result = self.classify(
            "Build warehouse dashboards for labor efficiency",
            "Built interactive dashboards for business performance and executive decisions",
        )
        self.assertEqual(result["classification"], "transferable")

    def test_weak_overlap_is_no_confirmed_evidence(self):
        result = self.classify(
            "Experience with ads recommendation ecommerce and search",
            "Built litigation analytics for legal case comparison",
        )
        self.assertEqual(result["classification"], "no_confirmed_evidence")

    def test_credentials_are_not_confirmed_by_project_stories(self):
        result = self.classify(
            "5+ years of data science experience",
            "Built data science models and analytics workflows",
        )
        self.assertEqual(result["classification"], "no_confirmed_evidence")

    def test_possessive_degree_wording_is_a_credential(self):
        result = self.classify(
            "Bachelor's degree in Statistics or Computer Science",
            "Applied statistics and computer science methods in an analytics project",
        )
        self.assertEqual(result["classification"], "no_confirmed_evidence")


class RetrievalCliTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.workspace = Path(self.temp.name)
        story = self.workspace / "Forecasting_Project_Complete story.docx"
        document = Document()
        document.add_heading("Action", level=1)
        document.add_paragraph("Built Python and SQL forecasting pipelines for product planning.")
        document.save(story)
        catalog = {
            "version": 1,
            "projects": [{
                "project_name": "Forecasting",
                "employer_or_experience": "Example",
                "source_document_path": story.name,
                "broad_project_themes": ["forecasting", "product analytics"],
                "relevant_role_families": ["data science"],
                "supporting_image_types": [],
            }],
        }
        agents = self.workspace / ".agents"
        agents.mkdir()
        (agents / "resume-tailor-project-catalog.json").write_text(json.dumps(catalog), encoding="utf-8")

    def tearDown(self):
        self.temp.cleanup()

    def run_script(self, *arguments):
        completed = subprocess.run(
            [sys.executable, str(SCRIPT), "--workspace", str(self.workspace), *arguments],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
        )
        return json.loads(completed.stdout)

    def test_requirement_mode_returns_per_requirement_candidates(self):
        payload = self.run_script(
            "--query",
            "Responsibilities\nBuild forecasting pipelines.\nPreferred Qualifications\nAdvertising experience is a plus.",
            "--format",
            "json",
        )

        self.assertEqual(payload["mode"], "requirements")
        self.assertEqual(len(payload["requirements"]), 2)
        self.assertEqual(payload["requirements"][0]["category"], "responsibility")
        self.assertEqual(payload["requirements"][0]["evidence_classification"], "direct")
        self.assertTrue(payload["requirements"][0]["suggested_candidates"])
        self.assertFalse(payload["requirements"][0]["blocks_resume_generation"])
        self.assertEqual(payload["requirements"][1]["category"], "preferred_qualification")
        self.assertEqual(
            payload["requirements"][1]["evidence_classification"],
            "no_confirmed_evidence",
        )
        self.assertEqual(payload["requirements"][1]["suggested_candidates"], [])
        self.assertIn("portfolio_score", payload["results"][0])

    def test_combined_inputs_receive_unique_requirement_ids(self):
        query_file = self.workspace / "jd.txt"
        query_file.write_text("Minimum Qualifications\nPython experience.", encoding="utf-8")
        payload = self.run_script(
            "--query",
            "Responsibilities\nBuild forecasting pipelines.",
            "--query-file",
            str(query_file),
            "--format",
            "json",
        )

        self.assertEqual(
            [requirement["requirement_id"] for requirement in payload["requirements"]],
            ["R1", "R2"],
        )

    def test_flat_mode_remains_available(self):
        payload = self.run_script(
            "--query", "Python SQL forecasting", "--mode", "flat", "--format", "json"
        )

        self.assertEqual(payload["mode"], "flat")
        self.assertEqual(payload["requirements"], [])
        self.assertTrue(payload["results"])


if __name__ == "__main__":
    unittest.main()
