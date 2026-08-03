import argparse
import json
from pathlib import Path

from docx import Document


def length(value):
    return None if value is None else round(value.inches, 3)


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract DOCX text and basic layout facts.")
    parser.add_argument("paths", nargs="+", type=Path)
    args = parser.parse_args()

    for path in args.paths:
        doc = Document(path)
        sections = [{
            "page_width_in": length(s.page_width),
            "page_height_in": length(s.page_height),
            "top_margin_in": length(s.top_margin),
            "bottom_margin_in": length(s.bottom_margin),
            "left_margin_in": length(s.left_margin),
            "right_margin_in": length(s.right_margin),
        } for s in doc.sections]
        paragraphs = []
        for i, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if text:
                paragraphs.append({"index": i, "style": p.style.name, "text": text})
        tables = []
        for ti, table in enumerate(doc.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append({"index": ti, "rows": rows})
        print(json.dumps({
            "path": str(path.resolve()),
            "sections": sections,
            "paragraphs": paragraphs,
            "tables": tables,
        }, ensure_ascii=False))


if __name__ == "__main__":
    main()
