import argparse
import hashlib
import json
import math
import os
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

from docx import Document

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


TOKEN_RE = re.compile(r"[a-z0-9][a-z0-9+#./-]*", re.I)
STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "but", "by", "for", "from", "has", "have",
    "in", "into", "is", "it", "of", "on", "or", "our", "such", "that", "the", "their", "this",
    "to", "using", "we", "will", "with", "you", "your"
}
EXPANSIONS = {
    "integrity": ["safety", "risk", "privacy", "compliance", "incident", "fraud", "violation", "monitoring"],
    "safety": ["integrity", "risk", "privacy", "compliance", "incident", "security", "moderation"],
    "risk": ["compliance", "privacy", "incident", "fraud", "anomaly", "exposure", "monitoring"],
    "experiment": ["testing", "evaluation", "benchmark", "uat", "comparison", "ab"],
    "experimentation": ["testing", "evaluation", "benchmark", "uat", "comparison", "ab"],
    "metrics": ["kpi", "monitoring", "dashboard", "measurement", "adoption", "conversion"],
    "product": ["adoption", "rollout", "funnel", "customer", "usage", "feature", "stakeholder"],
    "machine": ["model", "classification", "regression", "feature", "evaluation", "prediction"],
    "llm": ["genai", "rag", "retrieval", "prompt", "extraction", "responsible", "evaluation"],
    "privacy": ["incident", "compliance", "onetrust", "reporting", "regulatory", "risk"],
    "consumer": ["customer", "user", "behavior", "feedback", "interview", "survey", "segmentation", "adoption"],
    "research": ["study", "evaluation", "experiment", "testing", "analysis", "methodology"],
    "survey": ["questionnaire", "interview", "feedback", "qualitative", "quantitative", "consumer"],
    "qualitative": ["interview", "feedback", "research", "synthesis", "narrative"],
    "quantitative": ["statistical", "metrics", "analysis", "modeling", "measurement"],
    "pipeline": ["automation", "workflow", "processing", "data", "aws"],
    "dashboard": ["visualization", "kpi", "tableau", "power", "streamlit", "interactive"],
    "digital": ["automation", "platform", "workflow", "ai", "transformation"],
    "leadership": ["stakeholder", "cross-functional", "mentoring", "prioritization", "communication"],
}

DERIVATIVE_HEADINGS = (
    "skills demonstrated", "how i would explain", "interview-ready", "interview reference",
    "best resume bullet", "best way to package", "concise interview summary", "short resume project summary",
    "one-line summary", "short version"
)
PRIMARY_HEADINGS = (
    "action", "result", "task", "situation", "what i did", "business context", "analytical approach",
    "data pipeline", "feature engineering", "evaluation", "testing", "post-production", "post-launch"
)

REQUIREMENT_CATEGORY_WEIGHTS = {
    "responsibility": 1.35,
    "minimum_qualification": 1.25,
    "role_context": 1.10,
    "preferred_qualification": 0.80,
    "other": 1.00,
}
JD_HEADING_CATEGORIES = (
    (("minimum qualification", "required qualification", "requirements", "who we look for"), "minimum_qualification"),
    (("preferred qualification", "preferred skills", "nice to have"), "preferred_qualification"),
    (("responsibilities", "key responsibilities", "what the role entails", "what you'll do", "what you will do"), "responsibility"),
    (("about the team", "about the role", "job description", "role overview", "business unit"), "role_context"),
)


def tokens(text):
    values = [m.group(0).lower().strip("./-") for m in TOKEN_RE.finditer(text)]
    return [value for value in values if len(value) > 1 and value not in STOPWORDS]


def expanded_query(text):
    result = tokens(text)
    for token in list(result):
        result.extend(EXPANSIONS.get(token, []))
    return result


def normalized_heading(text):
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def heading_category(text):
    normalized = normalized_heading(text)
    for labels, category in JD_HEADING_CATEGORIES:
        if any(label in normalized for label in labels):
            return category
    return None


def split_requirement_text(text):
    text = re.sub(r"^[\s\-\u2022*]+", "", text).strip()
    if not text:
        return []
    # Preserve normal prose paragraphs. Split only explicit semicolon-delimited lists.
    parts = [part.strip() for part in re.split(r"\s*;\s+(?=[A-Z])", text) if part.strip()]
    return parts or [text]


def requirement_records_from_blocks(blocks):
    requirements = []
    has_category_headings = any(
        heading_category(text) is not None and (is_heading or len(tokens(text)) <= 8)
        for text, is_heading in blocks
    )
    category = None if has_category_headings else "other"
    for text, is_heading in blocks:
        text = text.strip()
        if not text:
            continue
        detected = heading_category(text) if is_heading or len(tokens(text)) <= 8 else None
        if is_heading or detected:
            if detected:
                category = detected
            if is_heading or len(tokens(text)) <= 8:
                continue
        if category is None:
            continue
        for statement in split_requirement_text(text):
            if len(tokens(statement)) < 2:
                continue
            requirements.append({
                "requirement_id": f"R{len(requirements) + 1}",
                "category": category,
                "weight": REQUIREMENT_CATEGORY_WEIGHTS[category],
                "text": statement,
            })
    return requirements


def requirements_from_docx(path: Path):
    document = Document(path)
    blocks = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = paragraph.style.name.lower() if paragraph.style else ""
        detected = heading_category(text)
        is_heading = style.startswith(("heading", "title")) or (detected is not None and len(tokens(text)) <= 8)
        blocks.append((text, is_heading))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        blocks.append((paragraph.text.strip(), False))
    return requirement_records_from_blocks(blocks)


def requirements_from_text(text):
    blocks = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        detected = heading_category(line)
        looks_like_heading = detected is not None and len(tokens(line)) <= 8
        blocks.append((line, looks_like_heading))
    return requirement_records_from_blocks(blocks)


def heading_weight(heading):
    lowered = heading.lower()
    if any(label in lowered for label in DERIVATIVE_HEADINGS):
        return 0.62, "derivative-summary"
    if "facts to confirm" in lowered or "limitation" in lowered:
        return 0.8, "caution-or-gap"
    if any(label in lowered for label in PRIMARY_HEADINGS):
        return 1.12, "primary-evidence"
    return 1.0, "standard"


def file_hash(path: Path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def extract_sections(path: Path):
    document = Document(path)
    sections = []
    current = {"heading": "Document opening", "paragraph_start": 0, "paragraph_end": 0, "paragraphs": []}
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if not text:
            continue
        is_heading = paragraph.style and paragraph.style.name.lower().startswith(("heading", "title"))
        if is_heading and current["paragraphs"]:
            current["paragraph_end"] = index - 1
            sections.append(current)
            current = {"heading": text, "paragraph_start": index, "paragraph_end": index, "paragraphs": []}
        elif is_heading:
            current["heading"] = text
            current["paragraph_start"] = index
        else:
            current["paragraphs"].append({"index": index, "text": text})
            current["paragraph_end"] = index
    if current["paragraphs"] or current["heading"] != "Document opening":
        sections.append(current)
    for index, section in enumerate(sections):
        section["section_index"] = index
        section["text"] = "\n".join(p["text"] for p in section.pop("paragraphs"))
    return sections


def cached_sections(path: Path, cache_dir: Path):
    digest = file_hash(path)
    cache_key = hashlib.sha256(str(path.resolve()).lower().encode("utf-8")).hexdigest()
    cache_file = cache_dir / f"{cache_key}.json"
    if cache_file.exists():
        cached = json.loads(cache_file.read_text(encoding="utf-8"))
        if cached.get("sha256") == digest:
            return cached["sections"], "hit", digest
    sections = extract_sections(path)
    cache_dir.mkdir(parents=True, exist_ok=True)
    cache_file.write_text(json.dumps({
        "source": str(path.resolve()), "sha256": digest, "sections": sections
    }, ensure_ascii=False), encoding="utf-8")
    return sections, "refreshed", digest


def read_catalog(root: Path, catalog_path: Path):
    if not catalog_path.exists():
        return []
    data = json.loads(catalog_path.read_text(encoding="utf-8"))
    records = []
    for item in data.get("projects", []):
        record = dict(item)
        record["absolute_path"] = (root / item["source_document_path"]).resolve()
        records.append(record)
    return records


def bm25_scores(doc_tokens, query_tokens):
    n = len(doc_tokens)
    avgdl = sum(len(doc) for doc in doc_tokens) / max(n, 1)
    dfs = Counter()
    for doc in doc_tokens:
        dfs.update(set(doc))
    scores = []
    query_counts = Counter(query_tokens)
    for doc in doc_tokens:
        counts = Counter(doc)
        score = 0.0
        for term, query_weight in query_counts.items():
            if term not in counts:
                continue
            idf = math.log(1 + (n - dfs[term] + 0.5) / (dfs[term] + 0.5))
            tf = counts[term]
            norm = tf + 1.5 * (1 - 0.75 + 0.75 * len(doc) / max(avgdl, 1))
            score += min(query_weight, 3) * idf * (tf * 2.5 / norm)
        scores.append(score)
    return scores


def verify(path: Path, section_index: int, context_sections: int):
    sections = extract_sections(path)
    if section_index < 0 or section_index >= len(sections):
        raise SystemExit(f"Section index out of range: 0..{len(sections)-1}")
    start = max(0, section_index - max(context_sections, 0))
    end = min(len(sections), section_index + max(context_sections, 0) + 1)
    section = sections[section_index]
    print(json.dumps({
        "status": "source-verification",
        "source_document": str(path.resolve()),
        "selected_section": section,
        "source_context": sections[start:end],
    }, ensure_ascii=False, indent=2))


def main():
    parser = argparse.ArgumentParser(description="Rank project-story sections without treating retrieval as confirmed evidence.")
    parser.add_argument("--workspace", type=Path, default=Path.cwd())
    parser.add_argument("--catalog", type=Path, default=Path(".agents/resume-tailor-project-catalog.json"))
    parser.add_argument("--query")
    parser.add_argument("--query-file", type=Path)
    parser.add_argument("--jd-docx", type=Path)
    parser.add_argument("--top", type=int, default=8)
    parser.add_argument("--per-project", type=int, default=3)
    parser.add_argument("--per-requirement", type=int, default=3)
    parser.add_argument("--mode", choices=["requirements", "flat"], default="requirements")
    parser.add_argument("--max-chars", type=int, default=1800)
    parser.add_argument("--format", choices=["json", "markdown"], default="markdown")
    parser.add_argument("--cache-dir", type=Path)
    parser.add_argument("--verify-path", type=Path)
    parser.add_argument("--section-index", type=int)
    parser.add_argument("--context-sections", type=int, default=0)
    args = parser.parse_args()

    root = args.workspace.resolve()
    if args.verify_path is not None:
        if args.section_index is None:
            raise SystemExit("--section-index is required with --verify-path")
        path = args.verify_path if args.verify_path.is_absolute() else root / args.verify_path
        verify(path, args.section_index, args.context_sections)
        return

    query_parts = [args.query or ""]
    requirements = requirements_from_text(args.query or "") if args.query else []
    if args.query_file:
        query_path = args.query_file if args.query_file.is_absolute() else root / args.query_file
        query_file_text = query_path.read_text(encoding="utf-8")
        query_parts.append(query_file_text)
        requirements.extend(requirements_from_text(query_file_text))
    if args.jd_docx:
        jd_path = args.jd_docx if args.jd_docx.is_absolute() else root / args.jd_docx
        jd = Document(jd_path)
        query_parts.extend(p.text for p in jd.paragraphs if p.text.strip())
        query_parts.extend(cell.text for table in jd.tables for row in table.rows for cell in row.cells)
        requirements.extend(requirements_from_docx(jd_path))
    for index, requirement in enumerate(requirements, 1):
        requirement["requirement_id"] = f"R{index}"
    query_text = "\n".join(query_parts).strip()
    if not query_text:
        raise SystemExit("Provide --query and/or --jd-docx")

    catalog_path = args.catalog if args.catalog.is_absolute() else root / args.catalog
    records = read_catalog(root, catalog_path)
    cache_dir = args.cache_dir or Path(tempfile.gettempdir()) / "resume-tailor-cache"
    candidates = []
    cache_status = {}
    for project in records:
        path = project["absolute_path"]
        if not path.exists():
            continue
        sections, status, digest = cached_sections(path, cache_dir)
        cache_status[str(path)] = {"status": status, "sha256": digest}
        catalog_text = " ".join(
            [project.get("project_name", ""), project.get("employer_or_experience", "")]
            + project.get("broad_project_themes", []) + project.get("relevant_role_families", [])
        )
        for section in sections:
            candidates.append({
                "project": project,
                "section": section,
                "ranking_text": catalog_text + " " + section["heading"] + " " + section["text"],
            })

    document_tokens = [tokens(item["ranking_text"]) for item in candidates]

    def result_for(item, score, requirement=None):
        project, section = item["project"], item["section"]
        excerpt = section["text"][:args.max_chars]
        result = {
            "retrieval_status": "Candidate only - verify against source before confirming",
            "score": round(score, 4),
            "project_name": project.get("project_name"),
            "employer_or_experience": project.get("employer_or_experience"),
            "source_document_path": project.get("source_document_path"),
            "section_index": section["section_index"],
            "heading": section["heading"],
            "section_type": heading_weight(section["heading"])[1],
            "paragraph_start": section["paragraph_start"],
            "paragraph_end": section["paragraph_end"],
            "excerpt": excerpt,
            "excerpt_truncated": len(section["text"]) > len(excerpt),
        }
        if requirement:
            result["requirement_id"] = requirement["requirement_id"]
            result["requirement_category"] = requirement["category"]
        return result

    def rank_for_query(text, weight=1.0, limit=None):
        raw_scores = bm25_scores(document_tokens, expanded_query(text))
        scored = []
        for item, raw_score in zip(candidates, raw_scores):
            heading_multiplier, _ = heading_weight(item["section"]["heading"])
            scored.append((item, raw_score * heading_multiplier * weight))
        ranked = []
        project_counts = Counter()
        for item, score in sorted(scored, key=lambda pair: pair[1], reverse=True):
            if score <= 0:
                continue
            project_key = item["project"].get("source_document_path", "")
            if project_counts[project_key] >= max(args.per_project, 1):
                continue
            ranked.append((item, score))
            project_counts[project_key] += 1
            if len(ranked) >= max(limit or args.top, 1):
                break
        return ranked

    requirement_results = []
    portfolio_by_section = {}
    if args.mode == "requirements" and requirements:
        for requirement in requirements:
            ranked = rank_for_query(requirement["text"], requirement["weight"], args.per_requirement)
            matches = [result_for(item, score, requirement) for item, score in ranked]
            requirement_results.append({**requirement, "candidates": matches})
            for rank, match in enumerate(matches, 1):
                key = (match["source_document_path"], match["section_index"])
                prior_requirements = portfolio_by_section.get(key, {}).get("matched_requirements", [])
                prior_portfolio_score = portfolio_by_section.get(key, {}).get("portfolio_score", 0.0)
                if key not in portfolio_by_section or match["score"] > portfolio_by_section[key]["score"]:
                    portfolio_by_section[key] = dict(match)
                    portfolio_by_section[key]["matched_requirements"] = list(prior_requirements)
                portfolio_by_section[key]["portfolio_score"] = round(
                    prior_portfolio_score + requirement["weight"] / rank, 4
                )
                portfolio_by_section[key].setdefault("matched_requirements", [])
                if requirement["requirement_id"] not in portfolio_by_section[key]["matched_requirements"]:
                    portfolio_by_section[key]["matched_requirements"].append(requirement["requirement_id"])
        results = sorted(
            portfolio_by_section.values(), key=lambda item: item["portfolio_score"], reverse=True
        )[:max(args.top, 1)]
    else:
        results = [result_for(item, score) for item, score in rank_for_query(query_text)]

    payload = {
        "notice": "Retrieval results are not confirmed evidence. Re-open the original source section before drafting claims.",
        "cache_directory": str(cache_dir.resolve()),
        "cache_status": cache_status,
        "mode": "requirements" if args.mode == "requirements" and requirements else "flat",
        "requirement_weights": REQUIREMENT_CATEGORY_WEIGHTS,
        "requirements": requirement_results,
        "results": results,
    }
    if args.format == "json":
        print(json.dumps(payload, ensure_ascii=False, indent=2))
    else:
        print(payload["notice"])
        if payload["mode"] == "requirements":
            for requirement in requirement_results:
                print(f"\n# {requirement['requirement_id']} [{requirement['category']}, weight {requirement['weight']}] {requirement['text']}")
                for index, result in enumerate(requirement["candidates"], 1):
                    print(f"\n## {index}. {result['project_name']} - {result['heading']} (score {result['score']})")
                    print(f"Source: {result['source_document_path']} | section {result['section_index']} | paragraphs {result['paragraph_start']}-{result['paragraph_end']}")
                    print(result["excerpt"])
            return
        for index, result in enumerate(results, 1):
            print(f"\n## {index}. {result['project_name']} — {result['heading']} (score {result['score']})")
            print(f"Source: {result['source_document_path']} | section {result['section_index']} | paragraphs {result['paragraph_start']}-{result['paragraph_end']}")
            print(result["excerpt"])


if __name__ == "__main__":
    main()
