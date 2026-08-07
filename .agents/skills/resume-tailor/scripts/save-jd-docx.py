import argparse
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def main() -> None:
    parser = argparse.ArgumentParser(description="Save an unabridged job description as DOCX.")
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--title", required=True)
    parser.add_argument("--company", required=True)
    args = parser.parse_args()

    if args.output.exists():
        raise SystemExit(f"Refusing to overwrite: {args.output}")

    text = args.input.read_text(encoding="utf-8")
    args.output.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(4)
    run = heading.add_run(args.title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)

    company = doc.add_paragraph()
    company.paragraph_format.space_after = Pt(10)
    run = company.add_run(args.company)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        if not line:
            continue
        is_heading = line in {
            "Responsibilities", "About the Team", "Qualifications",
            "Minimum Qualifications", "Preferred Qualifications",
            "Responsibilities - What You'll Do", "Workstreams & Tasks",
            "Data Analyst Development", "Required", "Preferred",
        }
        paragraph.paragraph_format.keep_together = True
        if is_heading:
            paragraph.paragraph_format.space_before = Pt(7)
            paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(line)
        run.bold = is_heading
        run.font.name = "Arial"
        run.font.size = Pt(11 if is_heading else 10.5)

    doc.save(args.output)
    print("Saved JD DOCX successfully.")


if __name__ == "__main__":
    main()
